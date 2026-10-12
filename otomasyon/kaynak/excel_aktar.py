"""Excel'e aktarma - eski Bet365Pro'nun ürettiği dosyalarla aynı biçim.

Tüm Veriler : MacBilgi tablosunun tamamı (skor sarı, açılış camgöbeği, kapanış turuncu)
Açılış / Kapanış Verileri : tarih, saat, sonuç analizleri ve yalnızca açılış (_A)
                            ya da kapanış (_K) oranları; sonuç hücreleri renklendirilir.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SARI, TURUNCU, YESIL, KIRMIZI, BEYAZ, SIYAH, CAMGOBEGI = (
    "CDC610", "EF981C", "32CD32", "B22222", "FFFFFF", "000000", "008080")
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

_stil_onbellek = {}


def _stil(bg, fg):
    k = (bg, fg)
    if k not in _stil_onbellek:
        _stil_onbellek[k] = (PatternFill("solid", fgColor="FF" + bg), Font(color="FF" + fg))
    return _stil_onbellek[k]


def _boya(ws, r, c, bg, fg):
    hucre = ws.cell(r, c)
    hucre.fill, hucre.font = _stil(bg, fg)


# ------------------------------------------------------- analiz fonksiyonları
def _skor(s):
    a, b = s.replace(" ", "").split("-")
    return int(a), int(b)


def kazanan(s):
    if not s:
        return ""
    e, d = _skor(s)
    return "1" if e > d else "2" if d > e else "0"


def kg(s):
    if not s:
        return ""
    e, d = _skor(s)
    return "VAR" if e > 0 and d > 0 else "YOK"


def alt_ust(s, n):
    if not s:
        return ""
    e, d = _skor(s)
    return "ÜST" if e + d >= n else "ALT"


def ev_alt_ust(s, n):
    if not s:
        return ""
    return "ÜST" if _skor(s)[0] >= n else "ALT"


def dep_alt_ust(s, n):
    if not s:
        return ""
    return "ÜST" if _skor(s)[1] >= n else "ALT"


ORAN_KOLONLARI = (
    ["MS1", "MS0", "MS2", "IY1", "IY0", "IY2", "IIY1", "IIY0", "IIY2",
     "CS10", "CS12", "CS02", "IYCS10", "IYCS12", "IYCS02",
     ("KGVAR", "KGVAR"), ("KGYOK", "KGYOK"), ("IYKGVAR", "IY_KGVAR"), ("IYKGYOK", "IY_KGYOK"),
     ("IIYKGVAR", "IIY_KGVAR"), ("IIYKGYOK", "IIY_KGYOK")]
    + [f"{y}{n}5" for n in range(9) for y in ("ALT", "UST")]
    + [f"IY{y}{n}5" for n in range(5) for y in ("ALT", "UST")]
    + [(f"IYMS_{a}E{b}", f"IYMS_{a}E{b}") for a in "102" for b in "102"]
    + [(f"SCORE_{s}", f"Score_?{s}") for s in
       "10 20 30 40 50 60 70 21 31 41 51 61 32 42 52 43 00 11 22 33 01 02 03 04 05 06 07 "
       "12 13 14 15 16 23 24 25 34".split()]
    + [(f"IYSCORE_{s}", f"IYScore_?{s}") for s in
       "10 20 30 40 21 31 00 11 22 01 02 03 04 12 13".split()]
)


def _oran_kolonu(tanim, ek):
    """(Excel başlığı, MacBilgi kolonu). ek: 'A' (açılış) ya da 'K' (kapanış)."""
    if isinstance(tanim, str):
        return tanim, f"{tanim}_{ek}"
    baslik, kaynak = tanim
    return baslik, kaynak.replace("?", ek) if "?" in kaynak else f"{kaynak}_{ek}"


ANALIZ_BASLIKLARI = ["Tarih", "Saat", "Lig", "Ev", "Dep", "IYSkor", "MSSkor", "Iy", "Ms", "IyMs",
                     "MsKg", "IyKg", "AU25", "AU35", "IyAU05", "IyAU15", "Ev05U", "Ev15U",
                     "Dep05U", "Dep15U", "IyEv05U", "IyDep05U"]


def _tablo_ve_bicim(ws, satir_say, kolon_say, genislikler):
    son = f"{get_column_letter(kolon_say)}{max(satir_say + 1, 2)}"
    t = Table(displayName="Table1", ref=f"A1:{son}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(t)
    orta = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=satir_say + 1, max_col=kolon_say):
        for c in row:
            c.alignment = orta
    for i, w in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 6), 45)


def _baslik(ad):
    # EPPlus başlıklardaki '_' karakterini boşluğa çeviriyordu
    return ad.replace("_", " ")


def _genislik(deger):
    return len(str(deger)) if deger is not None else 0


def eski_sira(kolonlar):
    """Eski Bet365Pro'nun kolon sırası: 2Y (IIY) oranları İY'nin, 2Y KG oranları İY KG'nin arkasında."""
    iiy = [k for k in kolonlar if k.startswith("IIY") and not k.startswith("IIY_KG")]
    iiy_kg = [k for k in kolonlar if k.startswith("IIY_KG")]
    sira = []
    for k in kolonlar:
        if k in iiy or k in iiy_kg:
            continue
        sira.append(k)
        if k == "IY2_K":
            sira += iiy
        elif k == "IY_KGYOK_K":
            sira += iiy_kg
    return sira


def tum_verileri_yaz(kolonlar, satirlar, yol):
    sira = eski_sira(kolonlar)
    yer = [kolonlar.index(k) for k in sira]
    kolonlar, satirlar = sira, [[s[i] for i in yer] for s in satirlar]
    wb = Workbook()
    ws = wb.active
    ws.title = "Tüm Veriler"
    basliklar = [_baslik(k) for k in kolonlar]
    ws.append(basliklar)
    genislik = [len(b) for b in basliklar]
    i_bitmis, i_tarih = kolonlar.index("BitmisMac"), kolonlar.index("Tarih")
    for s in satirlar:
        s = list(s)
        s[i_bitmis] = bool(s[i_bitmis])
        try:
            s[i_tarih] = datetime.strptime(s[i_tarih], "%Y-%m-%d %H:%M:%S")
        except (TypeError, ValueError):
            pass
        ws.append(s)
        for i, v in enumerate(s):
            genislik[i] = max(genislik[i], _genislik(v) if i != i_tarih else 10)
    n = len(satirlar)
    for r in range(2, n + 2):
        ws.cell(r, i_tarih + 1).number_format = "MM/dd/yyyy"
        _boya(ws, r, 8, SARI, SIYAH)
        _boya(ws, r, 9, SARI, SIYAH)
        for c in range(10, len(kolonlar) + 1):
            if c % 2 == 0:
                _boya(ws, r, c, CAMGOBEGI, BEYAZ)
            else:
                _boya(ws, r, c, TURUNCU, SIYAH)
    _tablo_ve_bicim(ws, n, len(kolonlar), genislik)
    wb.save(yol)


def acilis_kapanis_yaz(kolonlar, satirlar, yol, ek):
    """ek: 'A' açılış, 'K' kapanış."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Açılış Verileri" if ek == "A" else "Kapanış Verileri"
    oranlar = [_oran_kolonu(t, ek) for t in ORAN_KOLONLARI]
    basliklar = ANALIZ_BASLIKLARI + [_baslik(b) for b, _ in oranlar]
    ws.append(basliklar)
    genislik = [len(b) for b in basliklar]
    idx = {k: i for i, k in enumerate(kolonlar)}

    for s in satirlar:
        g = lambda k: s[idx[k]] if k in idx else None  # noqa: E731
        try:
            t = datetime.strptime(g("Tarih"), "%Y-%m-%d %H:%M:%S")
            tarih, saat = f"{t:%d.%m.%Y} {GUNLER[t.weekday()]}", f"{t:%H:%M}"
        except (TypeError, ValueError):
            tarih, saat = g("Tarih"), ""
        iy, ms = g("IYSkor") or "", g("MSSkor") or ""
        try:
            analiz = [kazanan(iy), kazanan(ms),
                      f"{kazanan(iy)}/{kazanan(ms)}" if iy and ms else "",
                      kg(ms), kg(iy), alt_ust(ms, 3), alt_ust(ms, 4), alt_ust(iy, 1), alt_ust(iy, 2),
                      ev_alt_ust(ms, 1), ev_alt_ust(ms, 2), dep_alt_ust(ms, 1), dep_alt_ust(ms, 2),
                      ev_alt_ust(iy, 1), dep_alt_ust(iy, 1)]
        except ValueError:
            analiz = [""] * 15
        satir = [tarih, saat, g("Lig"), g("Ev"), g("Dep"), iy, ms] + analiz + [g(k) for _, k in oranlar]
        ws.append(satir)
        for i, v in enumerate(satir):
            genislik[i] = max(genislik[i], _genislik(v))

    n = len(satirlar)
    for r in range(2, n + 2):
        if not ws.cell(r, 7).value:  # başlamayan maç: renk yok
            continue
        for c in range(8, 23):
            v = str(ws.cell(r, c).value or "")
            if c in (8, 9) or c == 10:
                son = v[-1:] if c == 10 else v
                if son == "1":
                    _boya(ws, r, c, CAMGOBEGI, BEYAZ)
                elif son == "2":
                    _boya(ws, r, c, TURUNCU, SIYAH)
                else:
                    _boya(ws, r, c, SIYAH, BEYAZ)
            elif c in (11, 12):
                _boya(ws, r, c, YESIL if v == "VAR" else KIRMIZI, SIYAH)
            else:
                _boya(ws, r, c, YESIL if v == "ÜST" else KIRMIZI, SIYAH)
    _tablo_ve_bicim(ws, n, len(basliklar), genislik)
    wb.save(yol)
