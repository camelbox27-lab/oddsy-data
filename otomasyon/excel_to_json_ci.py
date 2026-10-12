"""Excel -> JSON donusturucu (sunucu surumu).

Repodaki bet365/excel_to_json.py ve iddaa/gunluk/excel_to_json.py scriptleri
Windows'ta elle kullanilmak icin yazilmis: input() ile bekliyor ve kendileri git
push ediyor. Bu surum ayni JSON cikti bicimini uretir ama etkilesim ve git
islemi yapmaz (push'u Actions akisi yapar).

Cikti bicimi, sitenin okudugu mevcut JSON'larla birebir ayni olmak zorundadir;
bu yuzden ayni kutuphaneler ayni parametrelerle kullanilir.
"""
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent


def iddaa() -> int:
    """iddaa/gunluk/gunlukmaclar.xlsx -> gunlukmaclar.json"""
    import openpyxl

    klasor = KOK / "iddaa" / "gunluk"
    excel = klasor / "gunlukmaclar.xlsx"
    if not excel.exists():
        print(f"[json] HATA: {excel} yok", flush=True)
        return 1

    wb = openpyxl.load_workbook(excel)
    ws = wb.active
    basliklar = [h.value for h in ws[1]]

    satirlar = []
    for satir in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in satir):
            continue
        kayit = {}
        for anahtar, deger in zip(basliklar, satir):
            if deger is None:
                kayit[anahtar] = None
            elif isinstance(deger, (int, float)):
                kayit[anahtar] = deger
            else:
                kayit[anahtar] = str(deger)
        satirlar.append(kayit)

    hedef = klasor / "gunlukmaclar.json"
    with open(hedef, "w", encoding="utf-8") as f:
        json.dump(satirlar, f, ensure_ascii=False, indent=2)
    print(f"[json] iddaa OK -> {hedef.name} ({len(satirlar)} mac)", flush=True)
    return 0


def bet365(gun_str: str | None = None) -> int:
    """bet365/GGAAYYYY.xlsx -> GGAAYYYY.json (+ gunlukmaclar.json = en guncel gun)"""
    import pandas as pd

    klasor = KOK / "bet365"
    if gun_str:
        gun = datetime.strptime(gun_str, "%Y-%m-%d")
    else:
        gun = datetime.now()

    excel = klasor / f"{gun:%d%m%Y}.xlsx"
    if not excel.exists():
        print(f"[json] HATA: {excel} yok", flush=True)
        return 1

    df = pd.read_excel(excel)                       # ilk sayfa = "Tum Veriler"
    df["_gunTarihi"] = gun.strftime("%d.%m.%Y")
    hedef = klasor / f"{gun:%d%m%Y}.json"
    hedef.write_text(
        df.to_json(orient="records", force_ascii=False, indent=2), encoding="utf-8")
    print(f"[json] bet365 OK -> {hedef.name} ({len(df)} mac)", flush=True)

    # gunlukmaclar.json her zaman en guncel tarihli gunun kopyasi olur
    en_guncel, en_guncel_tarih = None, None
    for f in klasor.glob("*.json"):
        if len(f.name) != 13:                       # GGAAYYYY.json = 13 karakter
            continue
        try:
            t = datetime.strptime(f.stem, "%d%m%Y")
        except ValueError:
            continue
        if en_guncel_tarih is None or t > en_guncel_tarih:
            en_guncel, en_guncel_tarih = f, t

    if en_guncel:
        shutil.copy2(en_guncel, klasor / "gunlukmaclar.json")
        print(f"[json] gunlukmaclar.json = {en_guncel.name}", flush=True)
    return 0


if __name__ == "__main__":
    hangi = sys.argv[1] if len(sys.argv) > 1 else ""
    gun = sys.argv[2] if len(sys.argv) > 2 else None
    if hangi == "iddaa":
        sys.exit(iddaa())
    if hangi == "bet365":
        sys.exit(bet365(gun))
    print("kullanim: excel_to_json_ci.py {iddaa|bet365} [YYYY-MM-DD]")
    sys.exit(2)
